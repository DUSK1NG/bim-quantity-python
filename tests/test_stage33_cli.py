"""End-to-end contracts for the stage 3.3 pipeline and report CLIs."""

from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).parents[1]
INPUT = ROOT / "data" / "sample" / "sample_elements.csv"
MANUAL = ROOT / "data" / "sample" / "sample_manual_validation.csv"
CONFIG_DIR = ROOT / "configs"
RUN_SCRIPT = ROOT / "scripts" / "run_pipeline.py"
REPORT_SCRIPT = ROOT / "scripts" / "export_report.py"
CLEAN_SCRIPT = ROOT / "scripts" / "clean_sample_data.py"

SHEETS = (
    "项目概览",
    "全部构件明细",
    "分楼层工程量",
    "分构件工程量",
    "分材料工程量",
    "示例造价汇总",
    "数据质量问题",
    "人工复核结果",
)


def _run(script: Path, output_dir: Path, *, config_dir: Path = CONFIG_DIR) -> subprocess.CompletedProcess[str]:
    """Run one stage 3.3 CLI with explicit paths and captured UTF-8 output."""

    return subprocess.run(
        [
            sys.executable,
            str(script),
            "--input",
            str(INPUT),
            "--config-dir",
            str(config_dir),
            "--manual",
            str(MANUAL),
            "--output-dir",
            str(output_dir),
        ],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=False,
    )


def test_stage33_cli_help_exposes_shared_arguments() -> None:
    """Both dedicated scripts expose the same four documented arguments."""

    for script in (RUN_SCRIPT, REPORT_SCRIPT):
        completed = subprocess.run(
            [sys.executable, str(script), "--help"],
            cwd=ROOT,
            capture_output=True,
            text=True,
            check=False,
        )
        assert completed.returncode == 0, completed.stderr
        assert "--input" in completed.stdout
        assert "--config-dir" in completed.stdout
        assert "--manual" in completed.stdout
        assert "--output-dir" in completed.stdout
        assert "Pipeline" in completed.stdout or "报表" in completed.stdout


def test_run_pipeline_cli_is_repeatable_and_writes_readable_intermediates(
    tmp_path: Path,
) -> None:
    """Two custom output directories receive byte-stable pipeline artifacts."""

    first_dir = tmp_path / "first" / "pipeline"
    second_dir = tmp_path / "second" / "pipeline"
    first = _run(RUN_SCRIPT, first_dir)
    second = _run(RUN_SCRIPT, second_dir)

    assert first.returncode == 2, first.stderr
    assert second.returncode == 2, second.stderr
    assert "240" in first.stdout
    assert "问题行" in first.stdout
    assert "示例总价" in first.stdout

    names = (
        "sample_elements_standard.csv",
        "sample_elements_quality_report.json",
        "sample_elements_summary.csv",
    )
    first_files = [first_dir / name for name in names]
    second_files = [second_dir / name for name in names]
    assert all(path.is_file() for path in first_files + second_files)
    assert [path.read_bytes() for path in first_files] == [
        path.read_bytes() for path in second_files
    ]

    standard = pd.read_csv(first_files[0], encoding="utf-8-sig")
    assert standard.shape == (240, 22)
    assert standard["source_file"].eq("sample_elements.csv").all()

    quality = json.loads(first_files[1].read_text(encoding="utf-8"))
    assert quality["total_rows"] == 240
    assert quality["source_file"] == "sample_elements.csv"
    assert quality["disclaimer"] == "本项目单价为教学示例数据，不用于正式工程造价。"
    quality_text = first_files[1].read_text(encoding="utf-8")
    assert "C:\\Users\\" not in quality_text
    assert "C:/Users/" not in quality_text

    summary = pd.read_csv(first_files[2], encoding="utf-8-sig")
    assert not summary.empty
    assert "total_cost" in summary.columns


def test_export_report_cli_is_repeatable_and_writes_eight_sheet_workbook(
    tmp_path: Path,
) -> None:
    """The report CLI writes a readable workbook and all six CSV exports."""

    first_dir = tmp_path / "first" / "reports"
    second_dir = tmp_path / "second" / "reports"
    first = _run(REPORT_SCRIPT, first_dir)
    second = _run(REPORT_SCRIPT, second_dir)

    assert first.returncode == 2, first.stderr
    assert second.returncode == 2, second.stderr
    workbook_path = first_dir / "sample_elements_report.xlsx"
    second_workbook_path = second_dir / "sample_elements_report.xlsx"
    assert workbook_path.is_file()
    assert second_workbook_path.is_file()

    workbook = load_workbook(workbook_path, read_only=False)
    assert tuple(workbook.sheetnames) == SHEETS
    assert workbook["全部构件明细"].max_row == 241
    assert workbook["数据质量问题"].max_row > 1
    overview_values = [
        cell.value
        for row in workbook["项目概览"].iter_rows()
        for cell in row
    ]
    assert "sample_elements.csv" in overview_values
    assert "本项目单价为教学示例数据，不用于正式工程造价。" in overview_values

    csv_names = (
        "sample_elements_details.csv",
        "sample_elements_by_level.csv",
        "sample_elements_by_category.csv",
        "sample_elements_by_material.csv",
        "sample_elements_cost_summary.csv",
        "sample_elements_quality_issues.csv",
    )
    first_csv = [first_dir / name for name in csv_names]
    second_csv = [second_dir / name for name in csv_names]
    assert all(path.is_file() for path in first_csv + second_csv)
    assert [path.read_bytes() for path in first_csv] == [
        path.read_bytes() for path in second_csv
    ]
    details = pd.read_csv(first_csv[0], encoding="utf-8-sig")
    assert details.shape == (240, 22)
    assert details["source_file"].eq("sample_elements.csv").all()


def test_stage33_clis_return_one_with_chinese_repair_hint_for_bad_input(
    tmp_path: Path,
) -> None:
    """Missing input is a fatal CLI error with a beginner-facing hint."""

    missing_input = tmp_path / "does-not-exist.csv"
    output_dir = tmp_path / "bad-output"
    for script in (RUN_SCRIPT, REPORT_SCRIPT):
        completed = subprocess.run(
            [
                sys.executable,
                str(script),
                "--input",
                str(missing_input),
                "--config-dir",
                str(CONFIG_DIR),
                "--output-dir",
                str(output_dir),
            ],
            cwd=ROOT,
            capture_output=True,
            text=True,
            check=False,
        )
        assert completed.returncode == 1
        message = f"{completed.stdout}\n{completed.stderr}"
        assert "修复" in message or "请检查" in message


def test_stage33_clis_return_zero_when_quality_has_no_error(tmp_path: Path) -> None:
    """Severity configuration controls the 0-versus-2 completion status."""

    config_dir = tmp_path / "configs-no-error"
    shutil.copytree(CONFIG_DIR, config_dir)
    quality_path = config_dir / "quality_rules.json"
    quality = json.loads(quality_path.read_text(encoding="utf-8"))
    for rule in quality["rules"]:
        if rule.get("severity") == "Error":
            rule["severity"] = "Warning"
    quality_path.write_text(
        json.dumps(quality, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    for script in (RUN_SCRIPT, REPORT_SCRIPT):
        completed = _run(script, tmp_path / script.stem, config_dir=config_dir)
        assert completed.returncode == 0, completed.stderr


def test_stage31_clean_cli_still_accepts_the_shared_relative_contract(
    tmp_path: Path,
) -> None:
    """The stage 3.1 cleaning entry point remains a working regression path."""

    completed = subprocess.run(
        [
            sys.executable,
            str(CLEAN_SCRIPT),
            "--input",
            "data/sample/sample_elements.csv",
            "--output-dir",
            str(tmp_path / "processed"),
            "--report-dir",
            str(tmp_path / "reports"),
            "--config-dir",
            "configs",
        ],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=False,
    )
    assert completed.returncode == 0, completed.stderr
    assert (tmp_path / "processed" / "sample_elements_clean.csv").is_file()
    assert (tmp_path / "reports" / "sample_elements_quality_report.json").is_file()
