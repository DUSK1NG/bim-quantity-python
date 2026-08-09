"""Excel and CSV exports for :class:`src.pipeline.PipelineArtifacts`.

This module is deliberately an output layer.  It consumes the immutable
pipeline artifacts and does not recalculate quantities, costs, or quality
rules.  The workbook layout is kept stable so that a beginner can inspect it
with Excel or ``openpyxl`` and downstream scripts can rely on the sheet order.
"""

from __future__ import annotations

import os
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from src.pipeline import PipelineArtifacts
from src.quality_report import DISCLAIMER, quality_report_to_dict


SHEET_NAMES: tuple[str, ...] = (
    "项目概览",
    "全部构件明细",
    "分楼层工程量",
    "分构件工程量",
    "分材料工程量",
    "示例造价汇总",
    "数据质量问题",
    "人工复核结果",
)

_CSV_EXPORT_SPECS: tuple[tuple[str, str], ...] = (
    ("details", "全部构件明细"),
    ("by_level", "分楼层工程量"),
    ("by_category", "分构件工程量"),
    ("by_material", "分材料工程量"),
    ("cost_summary", "示例造价汇总"),
    ("quality_issues", "数据质量问题"),
)

_ENGINEERING_COLUMNS = frozenset(
    {
        "length_m",
        "area_m2",
        "volume_m3",
        "quantity",
        "quantity_sum",
        "area_sum",
        "volume_sum",
        "auto_quantity",
        "manual_quantity",
        "absolute_error",
    }
)
_AMOUNT_COLUMNS = frozenset({"unit_price", "total_cost"})
_PERCENT_COLUMNS = frozenset(
    {"relative_error_pct", "mean_relative_error_pct", "max_relative_error_pct"}
)

_QUALITY_COLUMNS: tuple[str, ...] = (
    "raw_row_number",
    "rule_id",
    "severity",
    "field",
    "message",
    "source_file",
    "guid",
    "element_name",
    "type_name",
    "level",
    "suggestion",
)
_VALIDATION_COLUMNS: tuple[str, ...] = (
    "match_key_type",
    "match_key",
    "element_id",
    "guid",
    "raw_row_number",
    "source_file",
    "ifc_class",
    "category",
    "level",
    "auto_quantity",
    "manual_quantity",
    "absolute_error",
    "relative_error_pct",
    "manual_row_number",
    "message",
)

_ERROR_FILL = PatternFill(fill_type="solid", fgColor="F4CCCC")
_WARNING_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")
_INFO_FILL = PatternFill(fill_type="solid", fgColor="D9EAD3")


def _basename(value: Any) -> Any:
    """Return a platform-neutral basename without exposing local paths."""

    if value is None:
        return None
    try:
        missing = pd.isna(value)
        if missing is pd.NA:
            return None
        # ``numpy.bool_`` (returned for numpy scalar values) is deliberately
        # handled through ``bool`` as well; array-like values are not valid
        # trace fields and fall through to their string representation.
        if not hasattr(missing, "__len__") and bool(missing):
            return None
    except (TypeError, ValueError):
        pass
    normalized = str(value).replace("\\", "/")
    return normalized.rsplit("/", 1)[-1]


def _require_artifacts(artifacts: PipelineArtifacts) -> None:
    """Fail early with a useful message when a caller passes another object."""

    if not isinstance(artifacts, PipelineArtifacts):
        raise TypeError("artifacts 必须是 PipelineArtifacts。")


def _safe_frame(frame: pd.DataFrame, *, columns: Iterable[str] | None = None) -> pd.DataFrame:
    """Copy one output frame and normalize trace paths to basenames."""

    if not isinstance(frame, pd.DataFrame):
        raise TypeError("报表数据必须是 pandas DataFrame。")
    result = frame.copy(deep=True)
    if columns is not None:
        ordered = list(columns)
        for column in ordered:
            if column not in result.columns:
                result[column] = pd.NA
        result = result.loc[:, ordered]
    if "source_file" in result.columns:
        result["source_file"] = result["source_file"].map(_basename)
    return result


def _overview_frame(artifacts: PipelineArtifacts) -> pd.DataFrame:
    """Convert the pipeline overview dictionary to a readable two-column table."""

    labels = {
        "source_file": "来源文件",
        "total_rows": "总行数",
        "element_count": "可信构件数",
        "category_count": "类别数",
        "level_count": "楼层数",
        "area_sum": "总面积(m²)",
        "volume_sum": "总体积(m³)",
        "issue_rows": "问题行数",
        "total_cost": "示例总价",
        "generated_at": "生成时间",
        "disclaimer": "免责声明",
    }
    overview = dict(artifacts.overview)
    rows: list[tuple[str, Any]] = []
    # Keep the source from the artifact as the authoritative trace value.  It
    # is reduced to a basename even when a caller constructed the fixture with
    # an absolute Windows path.
    overview["source_file"] = _basename(artifacts.source_file)
    for key, value in overview.items():
        if key == "source_file":
            value = _basename(value)
        elif key == "disclaimer":
            value = value or artifacts.disclaimer or DISCLAIMER
        rows.append((labels.get(key, key), value))
    if "generated_at" not in overview:
        rows.append((labels["generated_at"], datetime.now().isoformat(timespec="seconds")))
    if not any(key == "免责声明" for key, _ in rows):
        rows.append((labels["disclaimer"], artifacts.disclaimer or DISCLAIMER))
    return pd.DataFrame(rows, columns=("项目指标", "值"))


def _quality_frame(artifacts: PipelineArtifacts) -> pd.DataFrame:
    """Serialize quality issues without inventing a second set of rules."""

    payload = quality_report_to_dict(artifacts.quality_report, _basename(artifacts.source_file))
    issues = payload.get("issues", [])
    result = pd.DataFrame(issues)
    result = _safe_frame(result, columns=_QUALITY_COLUMNS)
    if "source_file" in result.columns:
        source = _basename(artifacts.source_file)
        result["source_file"] = result["source_file"].where(
            result["source_file"].notna(), source
        )
    return result


def _validation_frame(artifacts: PipelineArtifacts) -> pd.DataFrame:
    """Return validation details, preserving typed columns for an empty result."""

    details = artifacts.validation_result.details
    return _safe_frame(details, columns=_VALIDATION_COLUMNS)


def _workbook_frames(artifacts: PipelineArtifacts) -> tuple[tuple[str, pd.DataFrame], ...]:
    """Build the fixed eight workbook tables in contract order."""

    return (
        (SHEET_NAMES[0], _overview_frame(artifacts)),
        (SHEET_NAMES[1], _safe_frame(artifacts.standard_frame)),
        (SHEET_NAMES[2], _safe_frame(artifacts.by_level)),
        (SHEET_NAMES[3], _safe_frame(artifacts.by_category)),
        (SHEET_NAMES[4], _safe_frame(artifacts.by_material)),
        (SHEET_NAMES[5], _safe_frame(artifacts.cost_summary)),
        (SHEET_NAMES[6], _quality_frame(artifacts)),
        (SHEET_NAMES[7], _validation_frame(artifacts)),
    )


def _number_format(column_name: str) -> str | None:
    if column_name in _ENGINEERING_COLUMNS:
        return "0.000"
    if column_name in _AMOUNT_COLUMNS:
        return "0.00"
    if column_name in _PERCENT_COLUMNS:
        return "0.00"
    return None


def _iter_width_values(sheet: Any, column_index: int) -> Iterable[str]:
    for row in sheet.iter_rows(min_col=column_index, max_col=column_index):
        value = row[0].value
        if value is not None:
            yield str(value)


def _style_sheet(sheet: Any) -> None:
    """Apply common readability, filtering, and numeric-formatting rules."""

    sheet.freeze_panes = "A2"
    max_column = max(1, int(sheet.max_column))
    max_row = max(1, int(sheet.max_row))
    sheet.auto_filter.ref = f"A1:{get_column_letter(max_column)}{max_row}"

    headers = {
        int(cell.column): str(cell.value)
        for cell in sheet[1]
        if cell.value is not None
    }
    for column_index in range(1, max_column + 1):
        header = headers.get(column_index, "")
        width = max((len(value) for value in _iter_width_values(sheet, column_index)), default=0)
        # Keep very long issue messages from making the workbook unusably wide.
        sheet.column_dimensions[get_column_letter(column_index)].width = min(
            50, max(10, width + 2)
        )
        number_format = _number_format(header)
        if number_format is not None:
            for row in sheet.iter_rows(min_row=2, min_col=column_index, max_col=column_index):
                row[0].number_format = number_format


def _style_quality_sheet(sheet: Any) -> None:
    headers = {
        str(cell.value): int(cell.column)
        for cell in sheet[1]
        if cell.value is not None
    }
    severity_column = headers.get("severity")
    if severity_column is None:
        return
    for row in range(2, sheet.max_row + 1):
        value = sheet.cell(row=row, column=severity_column).value
        if value == "Error":
            fill = _ERROR_FILL
        elif value == "Warning":
            fill = _WARNING_FILL
        elif value == "Info":
            fill = _INFO_FILL
        else:
            continue
        for cell in sheet[row]:
            cell.fill = fill


def _temporary_path(parent: Path, suffix: str) -> Path:
    """Create a closed temporary path in the destination directory."""

    handle, name = tempfile.mkstemp(prefix=".report-", suffix=suffix, dir=str(parent))
    os.close(handle)
    return Path(name)


def write_excel_report(artifacts: PipelineArtifacts, output_path: Path) -> None:
    """Write the fixed eight-sheet workbook using an atomic replace."""

    _require_artifacts(artifacts)
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = _temporary_path(output.parent, ".xlsx")
    try:
        frames = _workbook_frames(artifacts)
        with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
            for sheet_name, frame in frames:
                frame.to_excel(writer, sheet_name=sheet_name, index=False)

        workbook = load_workbook(temporary)
        for sheet in workbook.worksheets:
            _style_sheet(sheet)
        _style_quality_sheet(workbook[SHEET_NAMES[6]])
        workbook.save(temporary)
        os.replace(temporary, output)
    finally:
        if temporary.exists():
            temporary.unlink()


def _csv_frames(artifacts: PipelineArtifacts) -> tuple[tuple[str, pd.DataFrame], ...]:
    frames = dict(_workbook_frames(artifacts))
    return tuple(
        (suffix, frames[sheet_name])
        for suffix, sheet_name in _CSV_EXPORT_SPECS
    )


def write_csv_exports(artifacts: PipelineArtifacts, output_dir: Path) -> tuple[Path, ...]:
    """Write six UTF-8-SIG CSV exports and return their paths in stable order."""

    _require_artifacts(artifacts)
    directory = Path(output_dir)
    directory.mkdir(parents=True, exist_ok=True)
    stem = Path(_basename(artifacts.source_file) or "report").stem
    outputs: list[Path] = []
    for suffix, frame in _csv_frames(artifacts):
        target = directory / f"{stem}_{suffix}.csv"
        temporary = _temporary_path(directory, ".csv")
        try:
            frame.to_csv(temporary, index=False, encoding="utf-8-sig", lineterminator="\n")
            os.replace(temporary, target)
        finally:
            if temporary.exists():
                temporary.unlink()
        outputs.append(target)
    return tuple(outputs)


__all__ = [
    "SHEET_NAMES",
    "write_csv_exports",
    "write_excel_report",
]
